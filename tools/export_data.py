"""
Exports the data the Pokedex/Calculator app needs out of pokemon_full_dex.xlsx
into small JSON files under app/data/.

Run this again any time the spreadsheet changes:
    python export_data.py

It reads raw values only (no formulas) from the reference sheets:
Base Stats, Abilities, Move List, Move Tier Guide, Type Chart, Movesets.
The four calculator sheets are NOT read here -- their formulas were ported
by hand into js/calc.js so the web app doesn't need Excel/openpyxl
at runtime, only at export time.

data/special_moves.json (dice/rules for fixed-damage, OHKO and variable-power
moves the sheet has no dice for) is hand-maintained and NOT touched here.
"""
import json
import sys
from pathlib import Path

import openpyxl

SRC = Path(__file__).resolve().parent.parent.parent / "pokemon_full_dex.xlsx"
OUT = Path(__file__).resolve().parent.parent / "data"
OUT.mkdir(exist_ok=True)

GAME_CODES = {
    "scarlet-violet": 1,
    "sword-shield": 2,
    "ultra-sun-ultra-moon": 3,
    "brilliant-diamond-shining-pearl": 4,
    "lets-go-pikachu-lets-go-eevee": 5,
    "legends-arceus": 6,
    "omega-ruby-alpha-sapphire": 7,
}
GAME_LABELS = {v: k for k, v in GAME_CODES.items()}

METHOD_CODES = {"Level-Up": "L", "TM": "T", "Egg": "E", "Tutor": "U"}
METHOD_LABELS = {v: k for k, v in METHOD_CODES.items()}


def log(msg):
    print(msg, flush=True)


def main():
    log(f"Loading {SRC} ...")
    wb = openpyxl.load_workbook(SRC, data_only=True)

    # ---------- Base Stats + Abilities -> pokemon.json ----------
    log("Reading Base Stats + Abilities ...")
    base_rows = list(wb["Base Stats"].iter_rows(min_row=2, values_only=True))
    abil_rows = {r[0]: r for r in wb["Abilities"].iter_rows(min_row=2, values_only=True) if r[0] is not None}

    pokemon = []
    dex_to_idx = {}
    for r in base_rows:
        if r[0] is None or r[1] is None:
            continue
        dex, name, t1, t2, hp, atk, dfn, spa, spd, spe, total, gen, catch = r[:13]
        a = abil_rows.get(dex, ())
        entry = {
            "dex": dex,
            "name": name,
            "type1": t1,
            "type2": t2,
            "hp": hp, "atk": atk, "def": dfn, "spa": spa, "spd": spd, "spe": spe,
            "total": total,
            "gen": gen,
            "catchRate": catch,
            "ability1": a[2] if len(a) > 2 else None,
            "ability1Effect": a[3] if len(a) > 3 else None,
            "ability2": a[4] if len(a) > 4 else None,
            "ability2Effect": a[5] if len(a) > 5 else None,
            "hiddenAbility": a[6] if len(a) > 6 else None,
            "hiddenAbilityEffect": a[7] if len(a) > 7 else None,
        }
        dex_to_idx[dex] = len(pokemon)
        pokemon.append(entry)
    log(f"  {len(pokemon)} Pokemon")

    # ---------- Move List -> moves.json ----------
    log("Reading Move List ...")
    move_rows = list(wb["Move List"].iter_rows(min_row=2, values_only=True))
    moves = []
    name_to_move_idx = {}
    for r in move_rows:
        if r[0] is None:
            continue
        name, typ, cat, power, acc, pp, tier, dice, die, effect = r[:10]
        name_to_move_idx[name] = len(moves)
        moves.append({
            "name": name, "type": typ, "category": cat, "power": power,
            "accuracy": acc, "pp": pp, "tier": tier,
            "diceCount": dice, "dieSize": die, "effect": effect,
        })
    log(f"  {len(moves)} moves")

    # ---------- Move Tier Guide ----------
    log("Reading Move Tier Guide ...")
    tier_rows = list(wb["Move Tier Guide"].iter_rows(min_row=1, values_only=True))
    tier_guide = []
    example_moves = []
    section = None
    for r in tier_rows:
        if r[0] == "Tier" and r[1] == "Dice":
            section = "table"
            continue
        if r[0] == "Tier" and r[1] == "Example 1":
            section = "examples"
            continue
        if r[0] == "Example Moves by Tier" or r[0] is None or (r[0] and str(r[0]).startswith("Note:")):
            if r[0] == "Example Moves by Tier":
                section = None
            continue
        if section == "table":
            tier, dice, count, size, mn, avg, mx, prange, nmoves = r[:9]
            tier_guide.append({
                "tier": tier, "dice": dice, "diceCount": count, "dieSize": size,
                "minRoll": mn, "avgRoll": avg, "maxRoll": mx,
                "powerRange": prange, "moveCount": nmoves,
            })
        elif section == "examples":
            tier = r[0]
            examples = [c for c in r[1:5] if c]
            example_moves.append({"tier": tier, "examples": examples})

    # ---------- Type Chart ----------
    log("Reading Type Chart ...")
    tc_ws = wb["Type Chart"]
    tc_rows = list(tc_ws.iter_rows(min_row=4, values_only=True))
    header_row = tc_rows[0]
    def_types = [c for c in header_row[1:] if c is not None]  # defending types (row[0] is corner label)
    type_chart = {"defTypes": def_types, "rows": {}}
    for r in tc_rows[1:]:
        if not r[0]:
            continue
        atk_type = r[0]
        vals = list(r[1:1 + len(def_types)])
        type_chart["rows"][atk_type] = vals

    # ---------- Movesets -> movesets.json (compact array form) ----------
    log("Reading Movesets (this is the big one, ~74k rows) ...")
    ms_rows = list(wb["Movesets"].iter_rows(min_row=2, values_only=True))

    # A ~100 newer-gen moves show up in movesets but aren't in Move List (no
    # tier/dice/power data -- same moves the spreadsheet's own calculators
    # can't select either, since their dropdown is Move List!A2:A809). Add
    # them as "stub" entries so the Pokedex can still show a Pokemon knows
    # them, while calculators keep filtering to real (non-stub) moves only.
    extra_names = sorted({r[4] for r in ms_rows if r[4] and r[4] not in name_to_move_idx})
    for nm in extra_names:
        name_to_move_idx[nm] = len(moves)
        moves.append({
            "name": nm, "type": None, "category": None, "power": None,
            "accuracy": None, "pp": None, "tier": None,
            "diceCount": None, "dieSize": None, "effect": None,
            "stub": True,
        })
    if extra_names:
        log(f"  +{len(extra_names)} moves known only from Movesets (no Move List stats) added as stubs")

    compact = []
    unknown_games = set()
    unknown_methods = set()
    for r in ms_rows:
        dex, name, method, level, move, game = r[:6]
        if dex is None or move is None:
            continue
        midx = name_to_move_idx.get(move)
        if midx is None:
            continue  # shouldn't happen now that stubs are added above
        mcode = METHOD_CODES.get(method)
        if mcode is None:
            unknown_methods.add(method)
            mcode = method
        gcode = GAME_CODES.get(game)
        if gcode is None:
            unknown_games.add(game)
            gcode = game
        compact.append([dex, midx, mcode, level, gcode])
    log(f"  {len(compact)} moveset rows")
    if unknown_methods:
        log(f"  WARNING unknown methods: {unknown_methods}")
    if unknown_games:
        log(f"  WARNING unknown games: {unknown_games}")

    # ---------- write files ----------
    def dump(name, obj):
        path = OUT / name
        with open(path, "w", encoding="utf-8") as f:
            json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
        log(f"  wrote {path} ({path.stat().st_size / 1024:.0f} KB)")

    dump("pokemon.json", pokemon)
    dump("moves.json", moves)
    dump("type_chart.json", type_chart)
    dump("tier_guide.json", {"table": tier_guide, "examples": example_moves})
    dump("movesets.json", compact)
    dump("meta.json", {
        "gameLabels": GAME_LABELS,
        "methodLabels": METHOD_LABELS,
    })

    log("Done.")


if __name__ == "__main__":
    try:
        main()
    except FileNotFoundError as e:
        log(f"ERROR: {e}")
        sys.exit(1)
